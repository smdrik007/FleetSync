import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, header_bg="1E3A8A", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        if i == 0:
            # Repeat header on new pages
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for cell in row.cells:
                set_cell_background(cell, header_bg)
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
        else:
            bg = alt_bg if i % 2 == 1 else "FFFFFF"
            for cell in row.cells:
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=90, bottom=90, left=150, right=150)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.0)

def create_srs_document(output_path):
    doc = Document()

    # Page Margins: 1 inch all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("FleetSync SRS v2.0 | IEEE Std 830-1998 Compliant")
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Software Requirements Specification (SRS)")
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Navy #1E3A8A

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("FleetSync — SaaS-Based Multi-Tenant Vehicle Management & Real-Time Tracking System")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate 600

    # Metadata Card (Table)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document Version:", "2.0 (Final Release)"),
        ("Compliance Standard:", "IEEE Std 830-1998 Specification Standard"),
        ("Project Phase:", "Step 1 of 3: Software Requirements Specification (SRS)"),
        ("Target Platform:", "Multi-Tenant Web Application & Mobile Progressive Web App (PWA)")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        cell_k, cell_v = meta_table.rows[row_idx].cells
        cell_k.text = k
        cell_v.text = v
        set_cell_background(cell_k, "F8FAFC")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)
        cell_k.paragraphs[0].runs[0].font.bold = True
        cell_k.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell_v.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper functions for headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 118, 110) # Teal 700
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 65, 85)
        return h

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.bold = True
        p.add_run(text)
        return p

    # Revision History
    add_h2("Document Revision History")
    rev_table = doc.add_table(rows=5, cols=4)
    rev_headers = ["Version", "Date", "Author", "Summary of Changes"]
    for col_idx, h_text in enumerate(rev_headers):
        rev_table.rows[0].cells[col_idx].text = h_text
    
    rev_rows = [
        ("0.1", "Initial Draft", "Initial Author", "Initial single-school transport concept."),
        ("1.0", "Sep 14, 2026", "Architecture Team", "Expanded to SaaS multi-tenant platform with 4 roles."),
        ("1.1", "Sep 14, 2026", "Architecture Team", "Admin role removed; Transportation Pool consolidated; Route text box & Driver dropdown added; Driver privacy protection mandated."),
        ("2.0", "Sep 14, 2026", "Architecture Team", "Final Release: Integrated Smart Route Delimiter Parser and Graceful Signal Fallback ('Last Seen' badge); IEEE Std 830 finalization.")
    ]
    for r_idx, r_data in enumerate(rev_rows, start=1):
        for c_idx, val in enumerate(r_data):
            rev_table.rows[r_idx].cells[c_idx].text = val
    style_table(rev_table)

    # 1. Introduction
    add_h1("1. Introduction")
    
    add_h2("1.1 Purpose")
    doc.add_paragraph("This Software Requirements Specification (SRS) document details the complete functional, non-functional, external interface, and architectural requirements for FleetSync. It establishes the technical baseline agreed upon by stakeholders and developers, serving as the blueprint for System Architecture & Design (Step 2) and Feature-by-Feature Implementation (Step 3).")

    add_h2("1.2 Document Conventions")
    add_bullet(" Mandatory requirement for baseline deployment.", "MUST / SHALL:")
    add_bullet(" Highly recommended requirement to ensure usability and commercial viability.", "SHOULD:")
    add_bullet(" Optional enhancement planned for post-MVP releases.", "MAY:")
    add_bullet(" Format [FR-<MODULE>-<NUMBER>] for functional requirements and [NFR-<CATEGORY>-<NUMBER>] for non-functional requirements.", "Requirement IDs:")

    add_h2("1.3 Intended Audience")
    add_bullet(" To implement data models, real-time sync pipelines, and application interfaces.", "System Architects & Engineers:")
    add_bullet(" To inspect engineering rigor, requirement completeness, and project scope.", "Evaluators & Course Instructors:")
    add_bullet(" To verify operational fleet workflows, driver controls, and commuter views.", "End-User Stakeholders:")

    add_h2("1.4 Project Scope")
    doc.add_paragraph("FleetSync is a cloud-native, multi-tenant Software-as-a-Service (SaaS) platform providing unified fleet management, driver assignment, and live transit tracking. It eliminates dedicated expensive hardware GPS trackers by utilizing standard smart-device web geolocation (HTML5 Geolocation API) on drivers' smartphones.")
    doc.add_paragraph("The platform operates on a lean, three-role architecture:")
    add_bullet(" The exclusive administrative manager of the fleet. Sets routes via text input, assigns drivers via dropdown, updates vehicle activity status, and manages users.", "1. Transportation Pool (Authority):")
    add_bullet(" Operates the assigned vehicle, controls trip start/end cycles, and broadcasts live GPS location with strict privacy protections.", "2. Driver:")
    add_bullet(" Selects a vehicle number from a dropdown to view its real-time location on an interactive map, structured route stops, and assigned driver contact details.", "3. Passenger (Students, Staff, Commuters):")

    # 2. Overall Description
    add_h1("2. Overall Description")

    add_h2("2.1 Product Perspective")
    doc.add_paragraph("FleetSync is structured around a multi-tenant shared database architecture with logical tenant isolation enforced at the data-access layer via institutionId. Each institution operates within its secure logical partition.")

    add_h2("2.2 User Classes & Responsibilities")
    role_table = doc.add_table(rows=4, cols=3)
    role_headers = ["User Role", "Primary Responsibilities", "Core Screen Features"]
    for c_idx, h_text in enumerate(role_headers):
        role_table.rows[0].cells[c_idx].text = h_text
    role_rows = [
        ("Transportation Pool (Authority)", "Exclusive vehicle management, route setting, driver assignment, status control, fleet monitoring.", "Fleet roster, Route Text Box, Driver Dropdown, Status selector (Active/Trip/Maint/Inactive)."),
        ("Driver", "Operates vehicle, triggers trip lifecycle (Start/End), broadcasts location while on duty.", "Assigned vehicle display, Start/End Trip buttons, Live location pulse indicator, Reassignment alert."),
        ("Passenger (Commuters)", "Tracks shuttle/bus, views route stops, communicates with driver.", "Vehicle Dropdown Selector, Live Interactive Leaflet Map, Milestone Stepper, One-Touch Call Button.")
    ]
    for r_idx, r_data in enumerate(role_rows, start=1):
        for c_idx, val in enumerate(r_data):
            role_table.rows[r_idx].cells[c_idx].text = val
    style_table(role_table)

    add_h2("2.3 Operating Environment")
    add_bullet(" Google Chrome 90+, Mozilla Firefox 88+, Apple Safari 14+, Microsoft Edge 90+. Mobile responsive on Android and iOS.", "Supported Browsers:")
    add_bullet(" Node.js LTS (v18+ / v20+) with Next.js App Router.", "Application Server:")
    add_bullet(" PostgreSQL with Prisma ORM.", "Database Engine:")
    add_bullet(" WebSockets / Server-Sent Events (SSE) for sub-second telemetry and instant push.", "Real-Time Layer:")

    # 3. Functional Requirements
    add_h1("3. System Features & Specific Functional Requirements")

    add_h2("3.1 Module 1: Multi-Tenant Authentication & Institutional Login")
    add_bullet(" The login interface shall provide an institution selection mechanism allowing users to select their institution from a dynamic search dropdown or institution slug before authenticating.", "[FR-AUTH-01] Institution Profile Selection:")
    add_bullet(" Users sign in with their registered email/username and password. Passwords must be hashed using bcrypt (rounds >= 10). Sessions are strictly bound to the selected institutionId.", "[FR-AUTH-02] Secure Credential Verification:")
    add_bullet(" Upon successful login, users are routed automatically to their role console: Authority -> /authority/fleet, Driver -> /driver/console, Passenger -> /passenger/track.", "[FR-AUTH-03] Role-Based Redirection:")
    add_bullet(" Every backend query and mutation must enforce tenant isolation (WHERE institutionId = session.institutionId). Cross-tenant access is rejected with HTTP 403 Forbidden.", "[FR-AUTH-04] Strict Tenant Sandboxing:")

    add_h2("3.2 Module 2: Transportation Pool / Authority (Exclusive Vehicle Management)")
    p_auth_note = doc.add_paragraph()
    r_an = p_auth_note.add_run("Important Specification: The Transportation Pool (Authority) is the sole entity authorized to create, update, and modify vehicle information.")
    r_an.font.bold = True
    r_an.font.color.rgb = RGBColor(180, 83, 9)

    add_bullet(" The Authority can view, register, edit, and decommission vehicles with fields: Vehicle Number (e.g., 'Bus #04'), License Plate, Model, Capacity, Status, Route Text, and Assigned Driver.", "[FR-POOL-01] Vehicle Roster CRUD:")
    add_bullet(" The Authority shall configure and update the vehicle route using an intuitive Text Box (e.g., 'Route 4: Mirpur 10 -> Kazipara -> Farmgate -> Campus').", "[FR-POOL-02] Route Configuration via Text Box:")
    add_bullet(" The Authority shall assign a driver to a vehicle using a Driver Dropdown Selector listing active institutional drivers. Assigning a driver automatically unbinds them from any prior vehicle to prevent double-booking.", "[FR-POOL-03] Driver Assignment via Dropdown:")
    add_bullet(" The Authority can set vehicle status to ACTIVE, ON_TRIP, MAINTENANCE, or INACTIVE.", "[FR-POOL-04] Vehicle Activity Status Control:")
    add_bullet(" Any modification made by the Authority to a vehicle's route, driver, or status shall be dispatched immediately via the real-time event bus. Driver and Passenger screens watching that vehicle update instantly without page reloads.", "[FR-POOL-05] Instant Real-Time Synchronization:")

    add_h2("3.3 Module 3: Driver Portal & Mandatory Privacy Protection")
    add_bullet(" Upon login, the driver's screen prominently displays their Assigned Vehicle Number, Plate Number, Route Text, and Status.", "[FR-DRV-01] Assigned Vehicle Display:")
    add_bullet(" Driver can tap 'Start Trip' to transition vehicle to ON_TRIP and initiate GPS broadcasting. Driver taps 'End Trip' at the terminus to complete the run and shut off GPS.", "[FR-DRV-02] Trip Lifecycle Control:")
    add_bullet(" (1) Location tracking is strictly bound to active trips (IN_PROGRESS). (2) Driver screen shows a prominent pulsing indicator: '🔴 Live Location Sharing is ACTIVE'. (3) The instant the driver taps 'End Trip', all geolocation listeners (clearWatch) terminate immediately. (4) Zero location tracking occurs while off-duty or logged out. (5) Automatic prompt to conclude trip if stationary at destination for >15 minutes.", "[FR-DRV-03] Mandatory Driver Privacy Protection (Core Mandate):")
    add_bullet(" If the Transportation Pool reassigns the driver to another vehicle or modifies the route, an instant in-app banner immediately notifies the driver.", "[FR-DRV-04] Instant Reassignment Alert:")

    add_h2("3.4 Module 4: Passenger Portal ('Where's My Bus / Shuttle')")
    add_bullet(" Passengers select their vehicle from a clean, responsive Vehicle Number Dropdown (e.g., 'Bus #04 [DHA-METRO-KA-11-2233] — Route 3'). Vehicles in MAINTENANCE or INACTIVE are clearly badged and disabled.", "[FR-PAS-01] Vehicle Selector Dropdown:")
    add_bullet(" Selecting a vehicle renders an interactive Leaflet/OpenStreetMap view centered on the vehicle's real-time position with live animated marker updates.", "[FR-PAS-02] Live Interactive Map Tracking:")
    add_bullet(" The web application automatically parses the raw route text from the Authority's text box (supporting '->', '-->', or ',' delimiters) into a sequential visual Milestone Timeline / Stepper, allowing commuters to easily follow the sequence of stops without requiring manual GPS waypoint mapping by administrators.", "[FR-PAS-03] Smart Route Delimiter Parser (Recommendation 1):")
    add_bullet(" If a driver enters a cellular dead zone or their phone battery depletes mid-trip, the map does not crash or disappear. The vehicle marker remains visible at the last reported coordinate, turning amber with an exact timestamp badge: '⚠️ Signal Lost — Last updated 3 minutes ago'. Once connection resumes, it returns to green LIVE status.", "[FR-PAS-04] Graceful Signal Fallback & 'Last Seen' Badge (Recommendation 3):")
    add_bullet(" Displays a dedicated card with Driver's Full Name, Assigned Vehicle Number, and Driver's Direct Phone Number with a clickable one-touch call button (tel: protocol) for immediate commuter assistance.", "[FR-PAS-05] Driver Information & One-Touch Direct Call:")

    # 4. External Interface Requirements
    add_h1("4. External Interface Requirements")
    add_h2("4.1 User Interfaces")
    doc.add_paragraph("Responsive mobile-first layout built with Tailwind CSS. Driver UI incorporates oversized touch targets (>= 56x56 px) with high sunlight contrast. Passenger tracking interface features a full-screen interactive Leaflet map.")
    add_h2("4.2 Hardware & Sensor Interfaces")
    doc.add_paragraph("Client-side HTML5 Geolocation API (navigator.geolocation) directly interfacing with onboard smartphone GPS receivers. No external OBD-II or dedicated telematics hardware is required.")
    add_h2("4.3 Communications Protocols")
    doc.add_paragraph("All communication is secured over HTTPS / WSS (TLS 1.3). Real-time telemetry and state changes are streamed via WebSockets or Server-Sent Events. Phone dialer is triggered via the universal tel: URI scheme.")

    # 5. Non-Functional Requirements
    add_h1("5. Non-Functional Requirements (NFRs)")
    nfr_table = doc.add_table(rows=7, cols=3)
    nfr_headers = ["ID", "Category", "Requirement Specification"]
    for c_idx, h_text in enumerate(nfr_headers):
        nfr_table.rows[0].cells[c_idx].text = h_text
    nfr_rows = [
        ("NFR-PERF-01", "Performance", "GPS position update latency from driver device to passenger screen shall not exceed 2.0 seconds over active cellular connections."),
        ("NFR-PERF-02", "Performance", "Instant sync latency: Authority changes to routes, drivers, or statuses must reflect across all connected clients in <= 1.0 second."),
        ("NFR-PERF-03", "Performance", "Initial passenger tracking screen must fully load and render within 1.5 seconds on standard 4G mobile connections."),
        ("NFR-SEC-01", "Security", "Complete logical tenant isolation. Cross-institution data leakage is structurally impossible at the database query layer."),
        ("NFR-SEC-02", "Privacy", "Strict driver privacy enforcement. Geolocation data is collected only during active trips; off-duty tracking is strictly prohibited."),
        ("NFR-REL-01", "Reliability", "Automatic reconnection with exponential backoff (1s, 2s, 4s, up to 15s) upon temporary cellular dropouts.")
    ]
    for r_idx, r_data in enumerate(nfr_rows, start=1):
        for c_idx, val in enumerate(r_data):
            nfr_table.rows[r_idx].cells[c_idx].text = val
    style_table(nfr_table)

    # 6. Verification and Traceability Matrix
    add_h1("6. Verification and Traceability Matrix")
    matrix_table = doc.add_table(rows=11, cols=4)
    matrix_headers = ["Req ID", "Requirement Description", "Primary Stakeholder", "Verification Method"]
    for c_idx, h_text in enumerate(matrix_headers):
        matrix_table.rows[0].cells[c_idx].text = h_text
    matrix_rows = [
        ("FR-AUTH-01/02", "Institution login & tenant-scoped auth", "All Users", "Integration test with multiple test institutions"),
        ("FR-POOL-01", "Vehicle roster CRUD operations", "Authority", "Functional UI validation"),
        ("FR-POOL-02", "Route setting via Text Box", "Authority", "Input validation & delimiter test"),
        ("FR-POOL-03", "Driver assignment via Dropdown", "Authority", "Dropdown selection & double-booking prevention test"),
        ("FR-POOL-04", "Vehicle activity status toggle", "Authority", "State machine transition test"),
        ("FR-POOL-05", "Instant real-time synchronization", "Authority -> All", "Multi-client WebSocket latency verification"),
        ("FR-DRV-01/02", "Driver console & Trip Start/End", "Driver", "Lifecycle transition test"),
        ("FR-DRV-03", "Mandatory Driver Privacy Protection", "Driver", "Geolocation watch lifecycle code audit"),
        ("FR-PAS-01/02", "Vehicle dropdown & live Leaflet map", "Passenger", "End-to-end browser map test"),
        ("FR-PAS-03/04", "Delimiter parser & 'Last Seen' fallback", "Passenger", "Network drop & regex delimiter test")
    ]
    for r_idx, r_data in enumerate(matrix_rows, start=1):
        for c_idx, val in enumerate(r_data):
            matrix_table.rows[r_idx].cells[c_idx].text = val
    style_table(matrix_table)

    # 7. Roadmap
    add_h1("7. Three-Step Project Execution Roadmap")
    doc.add_paragraph("The project strictly progresses through three sequential engineering phases:")
    add_bullet(" Complete IEEE Std 830-1998 document established; covers multi-tenancy, Authority exclusivity, driver privacy, route parser, and fallback. Saved in docs/SRS.md and published as docs/FleetSync_SRS_Final.docx.", "Step 1: Software Requirements Specification (SRS) [APPROVED & COMPLETED]:")
    add_bullet(" Multi-Tenant Relational Database Schema (Prisma ERD), Real-Time Event Architecture (WebSockets/SSE), REST/Server Action API Contracts, and Responsive UI Wireframes.", "Step 2: System Architecture & Detailed Design [NEXT PHASE]:")
    add_bullet(" Feature 1 (Auth & Multi-Tenancy) -> Feature 2 (Authority Vehicle Console & Dropdown) -> Feature 3 (Driver GPS Broadcaster & Privacy Guard) -> Feature 4 (Passenger Live Map, Stepper & Direct Call) -> Feature 5 (Real-time Event Engine & Production Polish).", "Step 3: Feature-by-Feature Implementation:")

    # Save
    doc.save(output_path)
    print(f"Successfully generated professional SRS document at: {output_path}")

if __name__ == "__main__":
    out_file = os.path.abspath("docs/FleetSync_SRS_Final.docx")
    create_srs_document(out_file)
